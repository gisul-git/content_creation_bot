"""
File extraction utility - Extracts text from various file formats.
Supports: PDF, DOCX, PPTX, TXT, and images (OCR).
"""
import os
import tempfile
from typing import Dict, Optional, Tuple
from pathlib import Path


class FileExtractor:
    """Extracts text content from various file formats."""
    
    def __init__(self):
        self.supported_extensions = {
            '.pdf': self._extract_pdf,
            '.docx': self._extract_docx,
            '.pptx': self._extract_pptx,
            '.txt': self._extract_txt,
            '.jpg': self._extract_image,
            '.jpeg': self._extract_image,
            '.png': self._extract_image,
        }
    
    def extract(self, file_path: str, file_name: str) -> Dict[str, any]:
        """
        Extract text content from a file.
        
        Args:
            file_path: Path to the uploaded file
            file_name: Original filename
            
        Returns:
            Dictionary with 'content', 'file_name', 'file_type', and 'success'
        """
        file_ext = Path(file_name).suffix.lower()
        
        if file_ext not in self.supported_extensions:
            return {
                'success': False,
                'error': f'Unsupported file type: {file_ext}',
                'content': '',
                'file_name': file_name,
                'file_type': file_ext
            }
        
        try:
            extractor_func = self.supported_extensions[file_ext]
            content = extractor_func(file_path)
            
            # Ensure we return non-empty content or indicate extraction issue
            if not content or not content.strip():
                return {
                    'success': False,
                    'error': f'No text content extracted from {file_name}. The file may be empty, corrupted, or image-based.',
                    'content': '',
                    'file_name': file_name,
                    'file_type': file_ext
                }
            
            return {
                'success': True,
                'content': content,
                'file_name': file_name,
                'file_type': file_ext,
                'error': None
            }
        except ImportError as e:
            # Handle missing dependencies with clear message
            return {
                'success': False,
                'error': f'Missing dependency: {str(e)}. Please install required packages: pip install -r requirements.txt',
                'content': '',
                'file_name': file_name,
                'file_type': file_ext
            }
        except Exception as e:
            return {
                'success': False,
                'error': f'Error extracting content: {str(e)}',
                'content': '',
                'file_name': file_name,
                'file_type': file_ext
            }
    
    def _extract_pdf(self, file_path: str) -> str:
        """Extract text from PDF file."""
        # Try pdfplumber first (better extraction quality)
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                text_parts = []
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
                return '\n\n'.join(text_parts) if text_parts else ""
        except ImportError:
            # Fallback to PyPDF2
            try:
                import PyPDF2
                text_parts = []
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page in pdf_reader.pages:
                        text = page.extract_text()
                        if text:
                            text_parts.append(text)
                return '\n\n'.join(text_parts) if text_parts else ""
            except ImportError:
                raise ImportError("Please install PyPDF2 or pdfplumber: pip install PyPDF2 pdfplumber")
    
    def _extract_docx(self, file_path: str) -> str:
        """Extract text from DOCX file."""
        try:
            from docx import Document
        except ImportError:
            raise ImportError("Please install python-docx: pip install python-docx")
        
        doc = Document(file_path)
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        return '\n\n'.join(text_parts)
    
    def _extract_pptx(self, file_path: str) -> str:
        """Extract text from PPTX file."""
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("Please install python-pptx: pip install python-pptx")
        
        prs = Presentation(file_path)
        text_parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = [f"Slide {slide_num}:"]
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
            if len(slide_text) > 1:
                text_parts.append('\n'.join(slide_text))
        return '\n\n'.join(text_parts)
    
    def _extract_txt(self, file_path: str) -> str:
        """Extract text from TXT file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            # Try with different encoding
            with open(file_path, 'r', encoding='latin-1') as file:
                return file.read()
    
    def _extract_image(self, file_path: str) -> str:
        """Extract text from image using OCR."""
        try:
            from PIL import Image
            import pytesseract
        except ImportError:
            raise ImportError("Please install pytesseract and Pillow: pip install pytesseract Pillow")
        
        try:
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
            return text
        except Exception as e:
            raise Exception(f"OCR failed: {str(e)}")


def get_file_extractor() -> FileExtractor:
    """Get file extractor instance."""
    return FileExtractor()

